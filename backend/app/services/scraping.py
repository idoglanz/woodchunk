"""Fetch a supplier page politely and reduce it to text an LLM can read."""

from __future__ import annotations

import asyncio
import csv
import io
import ipaddress
import socket
from urllib import robotparser
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from openpyxl import load_workbook

USER_AGENT = "WoodchunkCatalogBot/1.0 (+catalog import for a woodworking planner)"
MAX_CHARS = 60_000
MAX_BYTES = 3_000_000  # a page bigger than this isn't a price list we can read
MAX_REDIRECTS = 5
BLOCKED_SUFFIXES = (".internal", ".local", ".localhost", ".lan", ".home.arpa")


class FetchError(RuntimeError):
    pass


async def check_public(url: str) -> None:
    """Refuse addresses on this server's own network (SSRF): only public internet hosts may be fetched."""
    parts = urlparse(url)
    host = (parts.hostname or "").rstrip(".").lower()
    if parts.scheme not in ("http", "https") or not host:
        raise FetchError("Enter a web address that starts with http:// or https://")
    if host == "localhost" or host.endswith(BLOCKED_SUFFIXES):
        raise FetchError("That address isn't a public website.")
    try:
        infos = await asyncio.get_running_loop().getaddrinfo(host, parts.port or (443 if parts.scheme == "https" else 80),
                                                              type=socket.SOCK_STREAM)
    except OSError as e:
        raise FetchError("Couldn't find that website. Check the address.") from e
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if not ip.is_global or ip.is_multicast:
            raise FetchError("That address isn't a public website.")


async def safe_get(client: httpx.AsyncClient, url: str, timeout: float = 20) -> tuple[httpx.Response, bytes]:
    """GET a public URL, checking every redirect hop and stopping at MAX_BYTES."""
    for _ in range(MAX_REDIRECTS + 1):
        await check_public(url)
        async with client.stream("GET", url, timeout=timeout) as r:
            if r.is_redirect and r.headers.get("location"):
                url = urljoin(str(r.url), r.headers["location"])
                continue
            body = b""
            async for chunk in r.aiter_bytes():
                body += chunk
                if len(body) > MAX_BYTES:
                    raise FetchError("That page is too large to read. Paste the product list as text instead.")
            return r, body
    raise FetchError("That page redirects too many times.")


async def allowed_by_robots(client: httpx.AsyncClient, url: str) -> bool:
    parts = urlparse(url)
    robots_url = f"{parts.scheme}://{parts.netloc}/robots.txt"
    try:
        r, body = await safe_get(client, robots_url, timeout=8)
    except (httpx.HTTPError, FetchError):
        return True  # unreachable robots.txt: treat as allowed, like most crawlers
    if r.status_code >= 400:
        return True
    rp = robotparser.RobotFileParser()
    rp.parse(body.decode(r.encoding or "utf-8", errors="replace").splitlines())
    return rp.can_fetch(USER_AGENT, url)


def html_to_text(html: str, base_url: str = "") -> str:
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "noscript", "svg", "header", "footer", "nav", "form", "iframe"]):
        tag.decompose()
    for a in soup.find_all("a", href=True):
        if a.get_text(strip=True) and base_url:
            a.append(f" <{urljoin(base_url, a['href'])}>")
    for row in soup.find_all("tr"):
        cells = [c.get_text(" ", strip=True) for c in row.find_all(["td", "th"])]
        row.replace_with(" | ".join(cells) + "\n")
    text = soup.get_text("\n", strip=True)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    out, seen = [], set()
    for ln in lines:
        if ln not in seen or len(ln) > 40:
            out.append(ln)
        seen.add(ln)
    return "\n".join(out)[:MAX_CHARS]


async def fetch_page_text(url: str) -> tuple[str, str]:
    """Return (page title, page text)."""
    await check_public(url)
    async with httpx.AsyncClient(headers={"User-Agent": USER_AGENT}, follow_redirects=False, timeout=20) as client:
        if not await allowed_by_robots(client, url):
            raise FetchError("This site asks automated tools not to read that page (robots.txt). "
                             "Paste the product list as text instead.")
        try:
            r, body = await safe_get(client, url)
        except httpx.HTTPError as e:
            # Don't echo network details back: they'd tell a user what's reachable from the server.
            raise FetchError("Couldn't load that page. Check the address, or paste the product list as text.") from e
    if r.status_code >= 400:
        raise FetchError("That page couldn't be loaded (the site returned an error).")
    ctype = r.headers.get("content-type", "")
    if "html" not in ctype and "text" not in ctype:
        raise FetchError("That address is not a web page. Upload the file instead.")
    html = body.decode(r.encoding or "utf-8", errors="replace")
    soup = BeautifulSoup(html, "lxml")
    title = soup.title.get_text(strip=True) if soup.title else urlparse(url).netloc
    text = html_to_text(html, str(r.url))
    if len(text) < 200:
        raise FetchError("The page has almost no readable text (it probably loads products with JavaScript). "
                         "Copy the product list from your browser and paste it as text.")
    return title, text


def table_file_to_text(data: bytes, filename: str, max_rows: int = 400) -> str:
    name = filename.lower()
    rows: list[list[str]] = []
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows.append([f"# sheet: {ws.title}"])
            for r in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in r]
                if any(c.strip() for c in cells):
                    rows.append(cells)
                if len(rows) >= max_rows:
                    break
    else:
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096]) if text.strip() else csv.excel
        for r in csv.reader(io.StringIO(text), dialect):
            if any(c.strip() for c in r):
                rows.append(r)
            if len(rows) >= max_rows:
                break
    return "\n".join(" | ".join(c.strip() for c in r) for r in rows)[:MAX_CHARS]
